"""Standalone ETA report output."""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .common import (
    autosize,
    clean_number,
    unhide_workbook_columns,
    write_number_cell,
    write_text_cell,
)
from .dps_pp import week_label_for_date

ETA_OUTPUT_NAME = "ETA.xlsx"
ETA_SHEET_NAME = "Sheet1"
ETA_FIRST_PERIOD_COL = 5


@dataclass(frozen=True)
class EtaReportRow:
    model: str
    part_no: str
    po_remain: float
    vendor: str
    eta: Sequence[float]


def _period_text(period: Any, attribute: str) -> str:
    value = getattr(period, attribute, "")
    return "" if value is None else str(value).strip()


def _period_start(period: Any) -> dt.date | None:
    value = getattr(period, "start", None)
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def _write_period_header(ws, periods: Sequence[Any]) -> None:
    for offset, period in enumerate(periods, start=ETA_FIRST_PERIOD_COL):
        start = _period_start(period)
        if start is not None:
            write_text_cell(ws.cell(1, offset), _period_text(period, "header1") or start.strftime("%b").upper())
            write_text_cell(ws.cell(2, offset), _period_text(period, "header2") or week_label_for_date(start))
            write_text_cell(ws.cell(3, offset), _period_text(period, "header3") or start.strftime("%a"))
            cell = ws.cell(4, offset)
            cell.value = start
            cell.number_format = "m/d;@"
        else:
            write_text_cell(ws.cell(1, offset), _period_text(period, "header1"))
            write_text_cell(ws.cell(2, offset), _period_text(period, "header2"))
            write_text_cell(ws.cell(3, offset), _period_text(period, "header3"))
            write_text_cell(ws.cell(4, offset), _period_text(period, "header4") or _period_text(period, "label"))


def write_eta_report(output_path: Path, periods: Sequence[Any], rows: Sequence[EtaReportRow]) -> dict[str, int | Path]:
    wb = Workbook()
    ws = wb.active
    ws.title = ETA_SHEET_NAME

    for col, value in enumerate(("model", "Part No", "PO Remain", "Vendor"), start=1):
        write_text_cell(ws.cell(4, col), value)
    _write_period_header(ws, periods)

    for row_idx, row in enumerate(rows, start=5):
        write_text_cell(ws.cell(row_idx, 1), row.model)
        write_text_cell(ws.cell(row_idx, 2), row.part_no)
        write_number_cell(ws.cell(row_idx, 3), clean_number(row.po_remain))
        write_text_cell(ws.cell(row_idx, 4), row.vendor)
        for offset, value in enumerate(row.eta, start=ETA_FIRST_PERIOD_COL):
            if value:
                write_number_cell(ws.cell(row_idx, offset), clean_number(value))

    last_col = max(ETA_FIRST_PERIOD_COL + len(periods) - 1, 4)
    last_row = max(len(rows) + 4, 4)
    for header_row in range(1, 5):
        for col in range(1, last_col + 1):
            cell = ws.cell(header_row, col)
            if cell.value is not None:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

    ws.auto_filter.ref = f"A4:{get_column_letter(last_col)}{last_row}"
    ws.freeze_panes = None
    for row_idx in range(1, 5):
        ws.row_dimensions[row_idx].height = 12.75
    ws.column_dimensions["A"].width = 14.5
    ws.column_dimensions["B"].width = 13.5
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 12
    for col in range(ETA_FIRST_PERIOD_COL, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    autosize(ws, maximum=28)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    unhide_workbook_columns(wb)
    wb.save(output_path)
    return {
        "output": output_path,
        "rows": len(rows),
        "periods": len(periods),
    }
